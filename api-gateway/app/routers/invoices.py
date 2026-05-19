import io
import uuid
from datetime import date as date_cls, datetime, timezone
from typing import Optional

from fastapi import APIRouter, Body, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.invoice import Invoice, InvoiceStatus
from app.models.user import User, UserRole
from app.models.vendor import Vendor
from app.schemas.common import ApiResponse, PaginatedResponse, PaginationMeta
from app.schemas.invoice import InvoiceCreate, InvoiceOut
from app.services.auth_service import current_user, require_admin
from app.services.invoice_service import InvoiceService

router = APIRouter()


@router.get("", response_model=PaginatedResponse[InvoiceOut])
async def list_invoices(
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    department_id: Optional[str] = None,
    unassigned: bool = Query(False, description="부서 미배정 청구서만"),
    status: Optional[str] = None,
    vendor_id: Optional[str] = None,
    date_from: Optional[date_cls] = Query(None, description="발행일자 시작"),
    date_to: Optional[date_cls] = Query(None, description="발행일자 끝"),
    amount_min: Optional[float] = None,
    amount_max: Optional[float] = None,
    search: Optional[str] = Query(None, description="청구서번호 부분 일치"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    # member는 자기 부서로 강제
    if user.role != UserRole.admin:
        department_id = user.department_id
        unassigned = False
        if not department_id:
            return PaginatedResponse(
                success=True, data=[],
                meta=PaginationMeta(total=0, page=page, limit=limit, pages=0),
            )

    service = InvoiceService(db)
    items, total = await service.list_with_filter(
        page=page, limit=limit, department_id=department_id, unassigned=unassigned,
        status=status, vendor_id=vendor_id,
        date_from=date_from, date_to=date_to,
        amount_min=amount_min, amount_max=amount_max, search=search,
    )
    return PaginatedResponse(
        success=True,
        data=[InvoiceOut.model_validate(i) for i in items],
        meta=PaginationMeta(total=total, page=page, limit=limit, pages=(total + limit - 1) // limit),
    )


@router.get("/department-counts", response_model=ApiResponse[dict])
async def invoices_department_counts(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    service = InvoiceService(db)
    counts = await service.department_counts()
    if user.role != UserRole.admin:
        if not user.department_id:
            return ApiResponse(success=True, data={"total": 0, "unassigned": 0, "by_department": {}})
        my_count = counts["by_department"].get(user.department_id, 0)
        return ApiResponse(success=True, data={
            "total": int(my_count),
            "unassigned": 0,
            "by_department": {user.department_id: int(my_count)},
        })
    return ApiResponse(success=True, data=counts)


@router.get("/status-counts", response_model=ApiResponse[dict])
async def invoices_status_counts(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """탭 카운트용 — 상태별 청구서 수. member는 자기 부서 한정."""
    from sqlalchemy import func as _func
    q = select(Invoice.status, _func.count(Invoice.id)).group_by(Invoice.status)
    if user.role != UserRole.admin:
        if not user.department_id:
            return ApiResponse(success=True, data={"total": 0, "by_status": {}})
        q = q.where(Invoice.department_id == user.department_id)
    r = await db.execute(q)
    by_status = {row[0].value if hasattr(row[0], 'value') else str(row[0]): int(row[1]) for row in r.all()}
    total = sum(by_status.values())
    return ApiResponse(success=True, data={"total": total, "by_status": by_status})


@router.post("", response_model=ApiResponse[InvoiceOut], status_code=201)
async def create_invoice(body: InvoiceCreate, db: AsyncSession = Depends(get_db)):
    service = InvoiceService(db)
    invoice = await service.create(body)
    return ApiResponse(success=True, data=InvoiceOut.model_validate(invoice))


@router.get("/export-xlsx")
async def export_invoices_xlsx(
    department_id: Optional[str] = None,
    unassigned: bool = Query(False),
    status: Optional[str] = None,
    vendor_id: Optional[str] = None,
    date_from: Optional[date_cls] = Query(None),
    date_to: Optional[date_cls] = Query(None),
    amount_min: Optional[float] = None,
    amount_max: Optional[float] = None,
    search: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """현재 필터/검색 기준 청구서 목록을 xlsx로 다운로드 — 묶인 정산/거래 세부 포함.

    시트1: 청구서 목록
    시트2: 묶인 정산 (청구서별)
    시트3: 거래 세부 (정산별)
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from app.models.settlement import Settlement
    from app.models.transaction import Transaction
    from app.models.account import Department

    # member는 자기 부서로 강제
    if user.role != UserRole.admin:
        department_id = user.department_id
        unassigned = False
        if not department_id:
            wb = Workbook()
            wb.active.title = "청구서 목록"
            wb.active["A1"] = "조회 가능한 부서가 없습니다."
            buf = io.BytesIO(); wb.save(buf); buf.seek(0)
            return StreamingResponse(buf,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={"Content-Disposition": f'attachment; filename="invoices-empty.xlsx"'})

    service = InvoiceService(db)
    invoices, _total = await service.list_with_filter(
        page=1, limit=10000,
        department_id=department_id, unassigned=unassigned,
        status=status, vendor_id=vendor_id,
        date_from=date_from, date_to=date_to,
        amount_min=amount_min, amount_max=amount_max, search=search,
    )

    vendor_ids = list({i.vendor_id for i in invoices if i.vendor_id})
    dept_ids = list({i.department_id for i in invoices if i.department_id})

    vendor_map = {}
    if vendor_ids:
        v_r = await db.execute(select(Vendor).where(Vendor.id.in_(vendor_ids)))
        vendor_map = {v.id: v for v in v_r.scalars().all()}
    dept_map = {}
    if dept_ids:
        d_r = await db.execute(select(Department).where(Department.id.in_(dept_ids)))
        dept_map = {d.id: d for d in d_r.scalars().all()}

    # 청구서에 묶인 정산들
    invoice_ids = [i.id for i in invoices]
    s_r = await db.execute(select(Settlement).where(Settlement.invoice_id.in_(invoice_ids)))
    settlements_by_invoice: dict[str, list] = {}
    all_tx_ids: list[str] = []
    for s in s_r.scalars().all():
        settlements_by_invoice.setdefault(s.invoice_id, []).append(s)
        all_tx_ids.extend(s.matched_transactions or [])

    tx_map = {}
    if all_tx_ids:
        t_r = await db.execute(select(Transaction).where(Transaction.id.in_(all_tx_ids)))
        tx_map = {t.id: t for t in t_r.scalars().all()}

    # === xlsx 생성 ===
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    right_align = Alignment(horizontal="right")

    def style_header(ws, n):
        for col in range(1, n + 1):
            c = ws.cell(row=1, column=col)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")

    # 시트1: 청구서 목록
    ws1 = wb.active
    ws1.title = "청구서 목록"
    h1 = ["청구서번호", "공급자", "사업자번호", "대표자", "담당 부서",
          "발행일", "만기일", "결제완료일", "공급가액", "부가세", "청구액", "상태", "묶인 정산수", "은행", "계좌번호", "예금주"]
    for col, name in enumerate(h1, start=1):
        ws1.cell(row=1, column=col, value=name)
    style_header(ws1, len(h1))

    for ridx, inv in enumerate(invoices, start=2):
        v = vendor_map.get(inv.vendor_id)
        d = dept_map.get(inv.department_id) if inv.department_id else None
        bank = (v.bank_info or {}) if v else {}
        n_settlements = len(settlements_by_invoice.get(inv.id, []))
        row = [
            inv.invoice_no,
            v.name if v else "",
            v.business_registration_no if v else "",
            v.representative if v else "",
            d.name if d else "",
            inv.issue_date.isoformat(),
            inv.due_date.isoformat(),
            inv.paid_at.date().isoformat() if inv.paid_at else "",
            float(inv.subtotal),
            float(inv.tax_amount),
            float(inv.total_amount),
            inv.status.value,
            n_settlements,
            bank.get("bank_name", ""),
            bank.get("account_no", ""),
            bank.get("holder", ""),
        ]
        for col, val in enumerate(row, start=1):
            c = ws1.cell(row=ridx, column=col, value=val)
            if col in (9, 10, 11):
                c.number_format = "#,##0"
                c.alignment = right_align
            elif col == 13:
                c.alignment = right_align

    widths1 = [22, 24, 16, 12, 14, 12, 12, 14, 14, 12, 14, 12, 10, 12, 22, 14]
    for col, w in enumerate(widths1, start=1):
        ws1.column_dimensions[chr(64 + col)].width = w

    # 시트2: 묶인 정산
    ws2 = wb.create_sheet("묶인 정산")
    h2 = ["청구서번호", "정산번호", "공급자", "정산기간", "거래수", "정산금액", "정산상태"]
    for col, name in enumerate(h2, start=1):
        ws2.cell(row=1, column=col, value=name)
    style_header(ws2, len(h2))

    ridx = 2
    for inv in invoices:
        for s in settlements_by_invoice.get(inv.id, []):
            v = vendor_map.get(s.vendor_id)
            row = [
                inv.invoice_no,
                s.settlement_no,
                v.name if v else "",
                f"{s.period_start} ~ {s.period_end}",
                len(s.matched_transactions or []),
                float(s.total_amount),
                s.status.value,
            ]
            for col, val in enumerate(row, start=1):
                c = ws2.cell(row=ridx, column=col, value=val)
                if col in (5, 6):
                    c.number_format = "#,##0"
                    c.alignment = right_align
            ridx += 1

    widths2 = [22, 22, 24, 24, 8, 16, 12]
    for col, w in enumerate(widths2, start=1):
        ws2.column_dimensions[chr(64 + col)].width = w

    # 시트3: 거래 세부 (청구서 → 정산 → 거래)
    ws3 = wb.create_sheet("거래 세부")
    h3 = ["청구서번호", "정산번호", "공급자", "사업자번호",
          "거래일자", "승인번호", "품목명", "공급가액", "세액", "합계금액", "사전입금", "결제완료일"]
    for col, name in enumerate(h3, start=1):
        ws3.cell(row=1, column=col, value=name)
    style_header(ws3, len(h3))

    ridx = 2
    for inv in invoices:
        for s in settlements_by_invoice.get(inv.id, []):
            v = vendor_map.get(s.vendor_id)
            for tx_id in (s.matched_transactions or []):
                t = tx_map.get(tx_id)
                if not t:
                    continue
                row = [
                    inv.invoice_no,
                    s.settlement_no,
                    v.name if v else "",
                    v.business_registration_no if v else "",
                    t.transaction_date.isoformat(),
                    t.external_id or "",
                    t.description or "",
                    float(t.supply_amount) if t.supply_amount is not None else None,
                    float(t.tax_amount) if t.tax_amount is not None else None,
                    float(t.amount),
                    "Y" if t.is_prepaid else "",
                    t.paid_at.isoformat() if t.paid_at else "",
                ]
                for col, val in enumerate(row, start=1):
                    c = ws3.cell(row=ridx, column=col, value=val)
                    if col in (8, 9, 10):
                        c.number_format = "#,##0"
                        c.alignment = right_align
                ridx += 1

    widths3 = [22, 22, 24, 16, 14, 30, 36, 14, 14, 16, 10, 14]
    for col, w in enumerate(widths3, start=1):
        ws3.column_dimensions[chr(64 + col)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"invoices-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.get("/payment-sheet")
async def payment_sheet(
    ids: str = Query(..., description="콤마 구분 청구서 ID 리스트"),
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require_admin),
):
    """결제 매칭 문서(xlsx) 다운로드. 공급자/계좌번호/청구액 매칭표.

    재무팀이 은행 이체 시 사용. paid 상태로 변경하지는 않음(별도 bulk-pay 호출).
    /{invoice_id}와의 path 충돌을 피하기 위해 위쪽에 등록.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    invoice_ids = [s.strip() for s in ids.split(",") if s.strip()]
    if not invoice_ids:
        raise HTTPException(400, "청구서 ID가 없습니다.")

    inv_r = await db.execute(
        select(Invoice).where(
            Invoice.id.in_(invoice_ids),
            # 결제 대상만: 결제완료(paid)·취소(cancelled)는 매칭 문서에 포함하지 않음
            Invoice.status.in_([InvoiceStatus.draft, InvoiceStatus.sent, InvoiceStatus.overdue]),
        ).order_by(Invoice.due_date)
    )
    invoices = list(inv_r.scalars().all())
    if not invoices:
        raise HTTPException(
            status_code=404,
            detail="선택한 청구서 중 결제 대상이 없습니다 (paid/cancelled는 제외됨).",
        )

    vendor_ids = list({i.vendor_id for i in invoices})
    v_r = await db.execute(select(Vendor).where(Vendor.id.in_(vendor_ids)))
    vendor_map = {v.id: v for v in v_r.scalars().all()}

    wb = Workbook()
    ws = wb.active
    ws.title = "결제매칭"

    title_font = Font(bold=True, size=14)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    border = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    right_align = Alignment(horizontal="right")
    center_align = Alignment(horizontal="center")

    ws["A1"] = f"결제 매칭표 — {datetime.now().strftime('%Y-%m-%d')}"
    ws["A1"].font = title_font
    ws.merge_cells("A1:J1")
    ws.row_dimensions[1].height = 24

    headers = ["청구서번호", "공급자명", "사업자번호", "대표자",
               "은행", "계좌번호", "예금주",
               "공급가액", "부가세", "청구액(부가세포함)"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border

    holder_mismatch_fill = PatternFill("solid", fgColor="FEF3C7")  # 노란색 — 예금주가 대표자와 다를 때 시각 강조

    row = 4
    total_subtotal = 0.0
    total_tax = 0.0
    total_amount = 0.0
    for inv in invoices:
        v = vendor_map.get(inv.vendor_id)
        bank = (v.bank_info or {}) if v else {}
        holder = bank.get("holder", "")
        representative = (v.representative or "") if v else ""
        values = [
            inv.invoice_no,
            v.name if v else "(미상)",
            v.business_registration_no if v else "",
            representative,
            bank.get("bank_name", ""),
            bank.get("account_no", ""),
            holder,
            float(inv.subtotal),
            float(inv.tax_amount),
            float(inv.total_amount),
        ]
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = border
            if col_idx >= 8:
                cell.number_format = "#,##0"
                cell.alignment = right_align
        # 예금주와 대표자가 다르면 두 컬럼(대표자=4, 예금주=7)에 노란 하이라이트
        if representative and holder and representative.replace(" ", "") != holder.replace(" ", ""):
            ws.cell(row=row, column=4).fill = holder_mismatch_fill
            ws.cell(row=row, column=7).fill = holder_mismatch_fill
        total_subtotal += float(inv.subtotal)
        total_tax += float(inv.tax_amount)
        total_amount += float(inv.total_amount)
        row += 1

    ws.cell(row=row, column=1, value="합계").font = Font(bold=True)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    for col_idx, val in enumerate([total_subtotal, total_tax, total_amount], start=8):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = Font(bold=True)
        cell.number_format = "#,##0"
        cell.alignment = right_align
        cell.border = border

    widths = [22, 24, 16, 12, 12, 22, 16, 14, 12, 16]
    for col_idx, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"payment-sheet-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.get("/{invoice_id}", response_model=ApiResponse[InvoiceOut])
async def get_invoice(invoice_id: uuid.UUID, db: AsyncSession = Depends(get_db)):
    service = InvoiceService(db)
    invoice = await service.get(invoice_id)
    return ApiResponse(success=True, data=InvoiceOut.model_validate(invoice))


@router.post("/{invoice_id}/send", response_model=ApiResponse[InvoiceOut])
async def send_invoice(invoice_id: uuid.UUID, db: AsyncSession = Depends(get_db)):
    service = InvoiceService(db)
    invoice = await service.send(invoice_id)
    return ApiResponse(success=True, data=InvoiceOut.model_validate(invoice))


@router.patch("/{invoice_id}/paid", response_model=ApiResponse[InvoiceOut])
async def mark_paid(invoice_id: uuid.UUID, db: AsyncSession = Depends(get_db)):
    service = InvoiceService(db)
    invoice = await service.mark_paid(invoice_id)
    return ApiResponse(success=True, data=InvoiceOut.model_validate(invoice))


class BulkPayBody(BaseModel):
    invoice_ids: list[str]


class BulkSendBody(BaseModel):
    invoice_ids: list[str]


@router.post("/bulk-send", response_model=ApiResponse[dict])
async def bulk_send_invoices(
    body: BulkSendBody,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """선택된 청구서를 일괄 발송(sent) 처리. draft 상태만 대상.

    member는 자기 부서 청구서만 발송 가능.
    """
    if not body.invoice_ids:
        return ApiResponse(success=True, data={"updated": 0})

    q = select(Invoice).where(
        Invoice.id.in_(body.invoice_ids),
        Invoice.status == InvoiceStatus.draft,
    )
    if user.role != UserRole.admin:
        if not user.department_id:
            return ApiResponse(success=True, data={"updated": 0})
        q = q.where(Invoice.department_id == user.department_id)

    rows = await db.execute(q)
    allowed = list(rows.scalars().all())
    if not allowed:
        return ApiResponse(success=True, data={"updated": 0})

    now_naive = datetime.now(timezone.utc).replace(tzinfo=None)
    await db.execute(
        update(Invoice).where(Invoice.id.in_([i.id for i in allowed]))
        .values(status=InvoiceStatus.sent, sent_at=now_naive)
    )
    await db.commit()
    return ApiResponse(success=True, data={"updated": len(allowed)})


@router.delete("/{invoice_id}", response_model=ApiResponse[dict])
async def delete_invoice(
    invoice_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """청구서 삭제. 묶여있던 정산들은 다시 미발행(approved) 상태로 돌아간다.

    paid 상태인 청구서는 admin만 삭제 가능 (실수 방지).
    """
    from app.models.settlement import Settlement
    r = await db.execute(select(Invoice).where(Invoice.id == str(invoice_id)))
    inv = r.scalar_one_or_none()
    if not inv:
        raise HTTPException(status_code=404, detail="청구서를 찾을 수 없습니다.")
    if user.role != UserRole.admin and inv.department_id != user.department_id:
        raise HTTPException(status_code=403, detail="다른 부서의 청구서는 삭제할 수 없습니다.")
    if inv.status == InvoiceStatus.paid and user.role != UserRole.admin:
        raise HTTPException(status_code=403, detail="결제완료된 청구서는 admin만 삭제할 수 있습니다.")

    # 묶여있던 정산들의 invoice_id NULL + paid→approved 되돌림 (다시 미발행 상태)
    from app.models.settlement import SettlementStatus as _SS
    await db.execute(
        update(Settlement).where(Settlement.invoice_id == inv.id)
        .values(invoice_id=None, status=_SS.approved)
    )
    await db.delete(inv)
    await db.commit()
    return ApiResponse(success=True, data={"id": str(invoice_id)})


class PostponeBody(BaseModel):
    due_date: Optional[date_cls] = None


@router.patch("/{invoice_id}/postpone", response_model=ApiResponse[InvoiceOut])
async def postpone_invoice(
    invoice_id: uuid.UUID,
    body: PostponeBody = Body(default_factory=PostponeBody),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """청구서 만기일 연장 (재청구). due_date 미지정 시 기존 만기일에서 +30일.

    overdue 상태였다면 sent로 복귀. paid/cancelled는 변경 불가.
    """
    from datetime import timedelta
    r = await db.execute(select(Invoice).where(Invoice.id == str(invoice_id)))
    inv = r.scalar_one_or_none()
    if not inv:
        raise HTTPException(status_code=404, detail="청구서를 찾을 수 없습니다.")
    if user.role != UserRole.admin and inv.department_id != user.department_id:
        raise HTTPException(status_code=403, detail="다른 부서의 청구서는 재청구할 수 없습니다.")
    if inv.status in (InvoiceStatus.paid, InvoiceStatus.cancelled):
        raise HTTPException(
            status_code=409,
            detail=f"{inv.status.value} 상태 청구서는 재청구할 수 없습니다.",
        )

    inv.due_date = body.due_date if body.due_date else inv.due_date + timedelta(days=30)
    if inv.status == InvoiceStatus.overdue:
        inv.status = InvoiceStatus.sent
    await db.commit()
    await db.refresh(inv)
    return ApiResponse(success=True, data=InvoiceOut.model_validate(inv))


@router.patch("/{invoice_id}/cancel", response_model=ApiResponse[InvoiceOut])
async def cancel_invoice(
    invoice_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """청구서 반려/취소 — status=cancelled. paid는 변경 불가."""
    r = await db.execute(select(Invoice).where(Invoice.id == str(invoice_id)))
    inv = r.scalar_one_or_none()
    if not inv:
        raise HTTPException(status_code=404, detail="청구서를 찾을 수 없습니다.")
    if user.role != UserRole.admin and inv.department_id != user.department_id:
        raise HTTPException(status_code=403, detail="다른 부서의 청구서는 반려할 수 없습니다.")
    if inv.status == InvoiceStatus.paid:
        raise HTTPException(status_code=409, detail="결제완료된 청구서는 반려할 수 없습니다. 삭제하세요.")
    inv.status = InvoiceStatus.cancelled
    await db.commit()
    await db.refresh(inv)
    return ApiResponse(success=True, data=InvoiceOut.model_validate(inv))


@router.post("/bulk-pay", response_model=ApiResponse[dict])
async def bulk_pay_invoices(
    body: BulkPayBody,
    db: AsyncSession = Depends(get_db),
    _admin: User = Depends(require_admin),
):
    """선택된 청구서를 일괄 결제완료(paid) 처리. admin(재무팀)만 가능."""
    if not body.invoice_ids:
        return ApiResponse(success=True, data={"updated": 0})

    rows = await db.execute(
        select(Invoice).where(
            Invoice.id.in_(body.invoice_ids),
            Invoice.status.in_([InvoiceStatus.draft, InvoiceStatus.sent, InvoiceStatus.overdue]),
        )
    )
    allowed = list(rows.scalars().all())
    if not allowed:
        return ApiResponse(success=True, data={"updated": 0})

    now_naive = datetime.now(timezone.utc).replace(tzinfo=None)
    invoice_ids = [i.id for i in allowed]
    await db.execute(
        update(Invoice).where(Invoice.id.in_(invoice_ids))
        .values(status=InvoiceStatus.paid, paid_at=now_naive)
    )
    # 묶인 정산도 paid로 자동 전환
    from app.models.settlement import Settlement, SettlementStatus
    await db.execute(
        update(Settlement).where(Settlement.invoice_id.in_(invoice_ids))
        .values(status=SettlementStatus.paid)
    )
    await db.commit()
    return ApiResponse(success=True, data={"updated": len(allowed)})


