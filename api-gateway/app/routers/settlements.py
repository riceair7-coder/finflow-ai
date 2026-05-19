import io
import os
import re
import uuid
from datetime import date as date_cls, datetime
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.models.attachment import SettlementAttachment
from app.models.settlement import Settlement
from app.models.transaction import Transaction
from app.models.user import User, UserRole
from app.models.vendor import Vendor
from app.models.account import Department
from app.schemas.common import ApiResponse, PaginatedResponse, PaginationMeta
from app.schemas.invoice import InvoiceOut
from app.schemas.settlement import (
    IssueInvoiceFromSettlements,
    SettlementBulkCreate,
    SettlementCreate,
    SettlementOut,
    SettlementReview,
)
from app.services.auth_service import current_user
from app.services.invoice_service import InvoiceService
from app.services.settlement_service import SettlementService

router = APIRouter()


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# 첨부파일 저장 루트 — /app/storage/settlement-attachments/{settlement_id}/{att_id}_{safe_filename}
ATTACHMENT_ROOT = Path(os.environ.get("ATTACHMENT_ROOT", "/app/storage/settlement-attachments")).resolve()
ATTACHMENT_MAX_BYTES = 20 * 1024 * 1024  # 20MB
ATTACHMENT_ALLOWED_EXT = {
    ".pdf", ".png", ".jpg", ".jpeg", ".webp", ".gif",
    ".xlsx", ".xls", ".csv",
    ".doc", ".docx", ".hwp", ".hwpx",
    ".zip",
}
_SAFE_NAME_RE = re.compile(r"[^\w가-힣.\-\s()]+")


def _safe_filename(name: str) -> str:
    """경로 분리자/제어문자 제거 + 길이 제한."""
    base = os.path.basename(name).strip() or "file"
    base = _SAFE_NAME_RE.sub("_", base)
    return base[:200]


async def _load_settlement_for_user(db: AsyncSession, settlement_id: uuid.UUID, user: User) -> Settlement:
    row = await db.execute(select(Settlement).where(Settlement.id == str(settlement_id)))
    s = row.scalar_one_or_none()
    if not s:
        raise HTTPException(status_code=404, detail="정산을 찾을 수 없습니다.")
    if user.role != UserRole.admin and s.department_id != user.department_id:
        raise HTTPException(status_code=403, detail="다른 부서의 정산은 접근할 수 없습니다.")
    return s


def _attachment_to_dict(a: SettlementAttachment) -> dict:
    return {
        "id": a.id,
        "settlement_id": a.settlement_id,
        "filename": a.filename,
        "content_type": a.content_type,
        "file_size": int(a.file_size or 0),
        "kind": a.kind,
        "uploaded_by": a.uploaded_by,
        "uploaded_at": a.uploaded_at.isoformat() if a.uploaded_at else None,
    }


def _apply_xlsx_style(ws, n_cols: int):
    """엑셀 헤더 스타일 + 컬럼 너비 기본값."""
    from openpyxl.styles import Alignment, Font, PatternFill
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")


@router.post("", response_model=ApiResponse[SettlementOut], status_code=201)
async def create_settlement(body: SettlementCreate, db: AsyncSession = Depends(get_db)):
    service = SettlementService(db)
    settlement = await service.create(body)
    return ApiResponse(success=True, data=SettlementOut.model_validate(settlement))


@router.post("/bulk-create", response_model=ApiResponse[list[SettlementOut]], status_code=201)
async def bulk_create_settlements(
    body: SettlementBulkCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """여러 공급자에 대해 같은 기간의 정산을 일괄 생성."""
    if not body.vendor_ids:
        return ApiResponse(success=True, data=[])
    service = SettlementService(db)
    created = await service.bulk_create(
        vendor_ids=body.vendor_ids,
        period_start=body.period_start,
        period_end=body.period_end,
        notes=body.notes,
    )
    return ApiResponse(success=True, data=[SettlementOut.model_validate(s) for s in created])


@router.get("/unsettled-by-vendor", response_model=ApiResponse[list[dict]])
async def unsettled_by_vendor(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """공급자별 미정산 거래 잔액 (사전입금 제외)."""
    service = SettlementService(db)
    department_id = user.department_id if user.role != UserRole.admin else None
    items = await service.unsettled_by_vendor(department_id=department_id)
    return ApiResponse(success=True, data=items)


@router.get("/prepaid-by-vendor", response_model=ApiResponse[list[dict]])
async def prepaid_by_vendor(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """공급자별 사전입금 거래 집계 (정산관리에서 표시용)."""
    service = SettlementService(db)
    department_id = user.department_id if user.role != UserRole.admin else None
    items = await service.prepaid_by_vendor(department_id=department_id)
    return ApiResponse(success=True, data=items)


@router.get("", response_model=PaginatedResponse[SettlementOut])
async def list_settlements(
    page: int = Query(1, ge=1),
    limit: int = Query(20, ge=1, le=100),
    vendor_id: Optional[uuid.UUID] = None,
    status: Optional[str] = None,
    department_id: Optional[str] = None,
    unassigned: bool = Query(False, description="부서 미배정 정산만"),
    date_from: Optional[date_cls] = Query(None, description="정산기간 시작 (기간 겹침 매치)"),
    date_to: Optional[date_cls] = Query(None),
    amount_min: Optional[float] = None,
    amount_max: Optional[float] = None,
    search: Optional[str] = Query(None, description="정산번호 부분 일치"),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    # member는 자기 부서로 강제
    if user.role != UserRole.admin:
        department_id = user.department_id
        unassigned = False
        if not department_id:
            return PaginatedResponse(
                success=True, data=[],
                meta=PaginationMeta(total=0, page=page, limit=limit, pages=0),
            )

    service = SettlementService(db)
    items, total = await service.list(
        page=page, limit=limit, vendor_id=vendor_id, status=status,
        department_id=department_id, unassigned=unassigned,
        date_from=date_from, date_to=date_to,
        amount_min=amount_min, amount_max=amount_max, search=search,
    )
    return PaginatedResponse(
        success=True,
        data=[SettlementOut.model_validate(s) for s in items],
        meta=PaginationMeta(total=total, page=page, limit=limit, pages=(total + limit - 1) // limit),
    )


@router.get("/export-xlsx")
async def export_settlements_xlsx(
    vendor_id: Optional[uuid.UUID] = None,
    status: Optional[str] = None,
    department_id: Optional[str] = None,
    unassigned: bool = Query(False),
    date_from: Optional[date_cls] = None,
    date_to: Optional[date_cls] = None,
    amount_min: Optional[float] = None,
    amount_max: Optional[float] = None,
    search: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """현재 필터/검색 기준 정산 목록을 xlsx로 다운로드.

    시트1: 정산 목록 (정산번호/공급자/부서/기간/거래수/금액/상태/청구서)
    시트2: 거래 세부 (정산번호/거래일자/승인번호/품목명/공급가액/세액/합계)
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    # member는 본인 부서로 강제
    if user.role != UserRole.admin:
        department_id = user.department_id
        unassigned = False
        if not department_id:
            wb = Workbook()
            ws = wb.active
            ws.title = "정산 목록"
            ws["A1"] = "조회 가능한 부서가 없습니다."
            buf = io.BytesIO(); wb.save(buf); buf.seek(0)
            fname = f"settlements-empty-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
            return StreamingResponse(buf, media_type=XLSX_MIME,
                                     headers={"Content-Disposition": f'attachment; filename="{fname}"'})

    service = SettlementService(db)
    # limit 매우 크게 — 전체 다운로드
    settlements, _total = await service.list(
        page=1, limit=10000,
        vendor_id=vendor_id, status=status,
        department_id=department_id, unassigned=unassigned,
        date_from=date_from, date_to=date_to,
        amount_min=amount_min, amount_max=amount_max, search=search,
    )

    # vendor/department/invoice 정보 join용 조회
    vendor_ids = list({s.vendor_id for s in settlements if s.vendor_id})
    dept_ids = list({s.department_id for s in settlements if s.department_id})
    invoice_ids = list({s.invoice_id for s in settlements if s.invoice_id})

    vendor_map = {}
    if vendor_ids:
        v_r = await db.execute(select(Vendor).where(Vendor.id.in_(vendor_ids)))
        vendor_map = {v.id: v for v in v_r.scalars().all()}
    dept_map = {}
    if dept_ids:
        d_r = await db.execute(select(Department).where(Department.id.in_(dept_ids)))
        dept_map = {d.id: d for d in d_r.scalars().all()}
    from app.models.invoice import Invoice
    invoice_map = {}
    if invoice_ids:
        i_r = await db.execute(select(Invoice).where(Invoice.id.in_(invoice_ids)))
        invoice_map = {i.id: i for i in i_r.scalars().all()}

    # 거래 일괄 조회 (matched_transactions union)
    all_tx_ids: list[str] = []
    for s in settlements:
        all_tx_ids.extend(s.matched_transactions or [])
    tx_map = {}
    if all_tx_ids:
        t_r = await db.execute(select(Transaction).where(Transaction.id.in_(all_tx_ids)))
        tx_map = {t.id: t for t in t_r.scalars().all()}

    # === xlsx 생성 ===
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "정산 목록"
    headers1 = ["정산번호", "공급자", "사업자번호", "담당 부서", "정산기간",
                "거래수", "금액(부가세포함)", "상태", "청구서번호", "메모"]
    for col_idx, h in enumerate(headers1, start=1):
        ws1.cell(row=1, column=col_idx, value=h)
    _apply_xlsx_style(ws1, len(headers1))

    right_align = Alignment(horizontal="right")
    for idx, s in enumerate(settlements, start=2):
        v = vendor_map.get(s.vendor_id)
        d = dept_map.get(s.department_id) if s.department_id else None
        inv = invoice_map.get(s.invoice_id) if s.invoice_id else None
        row = [
            s.settlement_no,
            v.name if v else "",
            v.business_registration_no if v else "",
            d.name if d else "",
            f"{s.period_start} ~ {s.period_end}",
            len(s.matched_transactions or []),
            float(s.total_amount),
            s.status.value,
            inv.invoice_no if inv else "",
            s.notes or "",
        ]
        for col_idx, val in enumerate(row, start=1):
            cell = ws1.cell(row=idx, column=col_idx, value=val)
            if col_idx in (6, 7):
                cell.number_format = "#,##0"
                cell.alignment = right_align

    widths1 = [22, 24, 16, 14, 24, 8, 16, 12, 22, 30]
    for col_idx, w in enumerate(widths1, start=1):
        ws1.column_dimensions[chr(64 + col_idx)].width = w

    # 시트2: 거래 세부
    ws2 = wb.create_sheet("거래 세부")
    headers2 = ["정산번호", "공급자", "사업자번호", "거래일자", "승인번호", "품목명",
                "공급가액", "세액", "합계금액", "사전입금", "결제완료일"]
    for col_idx, h in enumerate(headers2, start=1):
        ws2.cell(row=1, column=col_idx, value=h)
    _apply_xlsx_style(ws2, len(headers2))

    row_idx = 2
    for s in settlements:
        v = vendor_map.get(s.vendor_id)
        for tx_id in (s.matched_transactions or []):
            t = tx_map.get(tx_id)
            if not t:
                continue
            row = [
                s.settlement_no,
                v.name if v else "",
                v.business_registration_no if v else "",
                t.transaction_date.isoformat(),
                t.external_id or "",
                t.description or "",
                float(t.supply_amount) if t.supply_amount is not None else None,
                float(t.tax_amount) if t.tax_amount is not None else None,
                float(t.amount),
                "Y" if t.is_prepaid else "",
                t.paid_at.isoformat() if t.paid_at else "",
            ]
            for col_idx, val in enumerate(row, start=1):
                cell = ws2.cell(row=row_idx, column=col_idx, value=val)
                if col_idx in (7, 8, 9):
                    cell.number_format = "#,##0"
                    cell.alignment = right_align
            row_idx += 1

    widths2 = [22, 24, 16, 14, 30, 36, 14, 14, 16, 10, 14]
    for col_idx, w in enumerate(widths2, start=1):
        ws2.column_dimensions[chr(64 + col_idx)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"settlements-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type=XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.get("/status-counts", response_model=ApiResponse[dict])
async def settlements_status_counts(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """탭 카운트용 — 정산 상태별 수. member는 자기 부서 한정."""
    from sqlalchemy import func as _func
    q = select(Settlement.status, _func.count(Settlement.id)).group_by(Settlement.status)
    if user.role != UserRole.admin:
        if not user.department_id:
            return ApiResponse(success=True, data={"total": 0, "by_status": {}})
        q = q.where(Settlement.department_id == user.department_id)
    r = await db.execute(q)
    by_status = {row[0].value if hasattr(row[0], 'value') else str(row[0]): int(row[1]) for row in r.all()}
    total = sum(by_status.values())
    return ApiResponse(success=True, data={"total": total, "by_status": by_status})


@router.get("/attachment-counts", response_model=ApiResponse[dict])
async def settlements_attachment_counts(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """정산별 첨부파일 건수 맵 — {settlement_id: count}.

    member는 본인 부서 정산만, admin은 전체.
    """
    from sqlalchemy import func as _func
    if user.role != UserRole.admin:
        if not user.department_id:
            return ApiResponse(success=True, data={})
        q = (
            select(SettlementAttachment.settlement_id, _func.count(SettlementAttachment.id))
            .join(Settlement, Settlement.id == SettlementAttachment.settlement_id)
            .where(Settlement.department_id == user.department_id)
            .group_by(SettlementAttachment.settlement_id)
        )
    else:
        q = (
            select(SettlementAttachment.settlement_id, _func.count(SettlementAttachment.id))
            .group_by(SettlementAttachment.settlement_id)
        )
    r = await db.execute(q)
    counts = {row[0]: int(row[1]) for row in r.all()}
    return ApiResponse(success=True, data=counts)


@router.get("/department-counts", response_model=ApiResponse[dict])
async def settlements_department_counts(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """탭 배지용 부서별 정산 수 (member는 자기 부서만)."""
    service = SettlementService(db)
    counts = await service.department_counts()
    if user.role != UserRole.admin:
        if not user.department_id:
            return ApiResponse(success=True, data={"total": 0, "unassigned": 0, "by_department": {}})
        my_count = counts["by_department"].get(user.department_id, 0)
        return ApiResponse(success=True, data={
            "total": int(my_count),
            "unassigned": 0,
            "by_department": {user.department_id: int(my_count)},
        })
    return ApiResponse(success=True, data=counts)


@router.get("/{settlement_id}", response_model=ApiResponse[SettlementOut])
async def get_settlement(settlement_id: uuid.UUID, db: AsyncSession = Depends(get_db)):
    service = SettlementService(db)
    settlement = await service.get(settlement_id)
    return ApiResponse(success=True, data=SettlementOut.model_validate(settlement))


@router.get("/{settlement_id}/details", response_model=ApiResponse[dict])
async def get_settlement_details(
    settlement_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """정산 상세: 매칭된 거래 목록 + 청구서 정보 + 같은 청구서에 묶인 다른 정산들."""
    from sqlalchemy import select as _select
    from fastapi import HTTPException
    from app.models.transaction import Transaction
    from app.models.invoice import Invoice
    from app.models.settlement import Settlement as _Settlement

    s_row = await db.execute(_select(_Settlement).where(_Settlement.id == str(settlement_id)))
    s = s_row.scalar_one_or_none()
    if not s:
        raise HTTPException(status_code=404, detail="정산을 찾을 수 없습니다.")
    # member 권한 가드
    if user.role != UserRole.admin and s.department_id != user.department_id:
        raise HTTPException(status_code=403, detail="다른 부서의 정산은 조회할 수 없습니다.")

    # 매칭된 거래
    tx_list = []
    if s.matched_transactions:
        tx_r = await db.execute(_select(Transaction).where(Transaction.id.in_(s.matched_transactions)))
        for t in tx_r.scalars().all():
            tx_list.append({
                "id": t.id,
                "transaction_date": t.transaction_date.isoformat(),
                "external_id": t.external_id,
                "description": t.description,
                "vendor_id": t.vendor_id,
                "amount": float(t.amount),
                "supply_amount": float(t.supply_amount) if t.supply_amount is not None else None,
                "tax_amount": float(t.tax_amount) if t.tax_amount is not None else None,
                "is_prepaid": t.is_prepaid,
            })
        tx_list.sort(key=lambda x: x["transaction_date"])

    # 청구서 및 같은 청구서에 묶인 다른 정산들
    invoice_info = None
    related_settlements: list[dict] = []
    if s.invoice_id:
        inv_r = await db.execute(_select(Invoice).where(Invoice.id == s.invoice_id))
        inv = inv_r.scalar_one_or_none()
        if inv:
            invoice_info = {
                "id": inv.id,
                "invoice_no": inv.invoice_no,
                "issue_date": inv.issue_date.isoformat(),
                "due_date": inv.due_date.isoformat(),
                "subtotal": float(inv.subtotal),
                "tax_amount": float(inv.tax_amount),
                "total_amount": float(inv.total_amount),
                "status": inv.status.value,
            }
            related_r = await db.execute(
                _select(_Settlement)
                .where(_Settlement.invoice_id == inv.id, _Settlement.id != s.id)
                .order_by(_Settlement.period_start)
            )
            for rs in related_r.scalars().all():
                related_settlements.append({
                    "id": rs.id,
                    "settlement_no": rs.settlement_no,
                    "period_start": rs.period_start.isoformat(),
                    "period_end": rs.period_end.isoformat(),
                    "total_amount": float(rs.total_amount),
                })

    return ApiResponse(success=True, data={
        "settlement": SettlementOut.model_validate(s).model_dump(mode="json"),
        "transactions": tx_list,
        "invoice": invoice_info,
        "related_settlements": related_settlements,
    })


@router.patch("/{settlement_id}/approve", response_model=ApiResponse[SettlementOut])
async def approve_settlement(
    settlement_id: uuid.UUID,
    body: SettlementReview,
    db: AsyncSession = Depends(get_db),
):
    service = SettlementService(db)
    settlement = await service.approve(settlement_id, body)
    return ApiResponse(success=True, data=SettlementOut.model_validate(settlement))


@router.patch("/{settlement_id}/reject", response_model=ApiResponse[SettlementOut])
async def reject_settlement(
    settlement_id: uuid.UUID,
    body: SettlementReview,
    db: AsyncSession = Depends(get_db),
):
    service = SettlementService(db)
    settlement = await service.reject(settlement_id, body)
    return ApiResponse(success=True, data=SettlementOut.model_validate(settlement))


@router.post("/{settlement_id}/issue-invoice", response_model=ApiResponse[InvoiceOut], status_code=201)
async def issue_invoice_from_settlement(
    settlement_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
):
    """승인된 정산 1건을 청구서로 발행한다."""
    service = InvoiceService(db)
    invoice = await service.create_from_settlement(settlement_id)
    return ApiResponse(success=True, data=InvoiceOut.model_validate(invoice))


@router.delete("/{settlement_id}", response_model=ApiResponse[dict])
async def delete_settlement(
    settlement_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """정산 삭제. 청구서가 발행된 정산은 삭제 불가(invoice_id IS NOT NULL).

    matched_transactions에 묶여있던 거래들은 정산 삭제 후 다시 미정산으로 돌아간다.
    """
    from sqlalchemy import select as _select
    from fastapi import HTTPException
    row = await db.execute(_select(Settlement).where(Settlement.id == str(settlement_id)))
    s = row.scalar_one_or_none()
    if not s:
        raise HTTPException(status_code=404, detail="정산을 찾을 수 없습니다.")
    if s.invoice_id:
        raise HTTPException(
            status_code=409,
            detail="청구서가 발행된 정산은 삭제할 수 없습니다. 청구서를 먼저 처리하세요.",
        )
    # member는 본인 부서 정산만 삭제 가능
    if user.role != UserRole.admin and s.department_id != user.department_id:
        raise HTTPException(status_code=403, detail="다른 부서의 정산은 삭제할 수 없습니다.")
    await db.delete(s)
    await db.commit()
    return ApiResponse(success=True, data={"id": str(settlement_id)})


@router.post("/issue-invoice", response_model=ApiResponse[InvoiceOut], status_code=201)
async def issue_invoice_from_settlements(
    body: IssueInvoiceFromSettlements,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """여러 정산을 한 청구서로 통합 발행한다.

    - 모든 정산이 같은 공급자여야 함
    - 모두 status=approved이고 미발행 상태여야 함
    """
    service = InvoiceService(db)
    invoice = await service.create_from_settlements(
        settlement_ids=body.settlement_ids,
        due_date_days=body.due_date_days,
        notes=body.notes,
    )
    return ApiResponse(success=True, data=InvoiceOut.model_validate(invoice))


# ====================================================================
# 첨부파일 (거래명세서, 영수증, 견적서 등)
# ====================================================================

@router.get("/{settlement_id}/attachments", response_model=ApiResponse[list[dict]])
async def list_attachments(
    settlement_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """정산에 첨부된 파일 목록."""
    await _load_settlement_for_user(db, settlement_id, user)
    rows = await db.execute(
        select(SettlementAttachment)
        .where(SettlementAttachment.settlement_id == str(settlement_id))
        .order_by(SettlementAttachment.uploaded_at.desc())
    )
    items = [_attachment_to_dict(a) for a in rows.scalars().all()]
    return ApiResponse(success=True, data=items)


@router.post("/{settlement_id}/attachments", response_model=ApiResponse[dict], status_code=201)
async def upload_attachment(
    settlement_id: uuid.UUID,
    file: UploadFile = File(...),
    kind: Optional[str] = Form(None),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """파일 업로드 — multipart/form-data.

    - file: 업로드 파일
    - kind: '거래명세서' | '영수증' | '견적서' | '기타' (선택)
    """
    await _load_settlement_for_user(db, settlement_id, user)

    orig_name = _safe_filename(file.filename or "file")
    ext = Path(orig_name).suffix.lower()
    if ext and ext not in ATTACHMENT_ALLOWED_EXT:
        raise HTTPException(
            status_code=415,
            detail=f"허용되지 않는 파일 형식입니다: {ext}. (허용: pdf, 이미지, xlsx, docx, hwp 등)",
        )

    # 사이즈 체크하면서 디스크에 기록 — 큰 파일 메모리 폭발 방지
    att_id = str(uuid.uuid4())
    dest_dir = ATTACHMENT_ROOT / str(settlement_id)
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest_path = dest_dir / f"{att_id}_{orig_name}"

    total = 0
    chunk_size = 1024 * 1024
    try:
        with dest_path.open("wb") as fp:
            while True:
                chunk = await file.read(chunk_size)
                if not chunk:
                    break
                total += len(chunk)
                if total > ATTACHMENT_MAX_BYTES:
                    fp.close()
                    dest_path.unlink(missing_ok=True)
                    raise HTTPException(
                        status_code=413,
                        detail=f"파일 크기가 너무 큽니다. 최대 {ATTACHMENT_MAX_BYTES // (1024 * 1024)}MB까지 업로드 가능합니다.",
                    )
                fp.write(chunk)
    except HTTPException:
        raise
    except Exception as e:
        dest_path.unlink(missing_ok=True)
        raise HTTPException(status_code=500, detail=f"파일 저장 실패: {e}")

    att = SettlementAttachment(
        id=att_id,
        settlement_id=str(settlement_id),
        filename=orig_name,
        stored_path=str(dest_path),
        content_type=file.content_type,
        file_size=total,
        kind=(kind or None),
        uploaded_by=user.id,
    )
    db.add(att)
    await db.commit()
    await db.refresh(att)
    return ApiResponse(success=True, data=_attachment_to_dict(att))


async def _resolve_attachment(
    db: AsyncSession, settlement_id: uuid.UUID, attachment_id: uuid.UUID, user: User,
) -> tuple[SettlementAttachment, Path]:
    await _load_settlement_for_user(db, settlement_id, user)
    row = await db.execute(
        select(SettlementAttachment).where(
            SettlementAttachment.id == str(attachment_id),
            SettlementAttachment.settlement_id == str(settlement_id),
        )
    )
    att = row.scalar_one_or_none()
    if not att:
        raise HTTPException(status_code=404, detail="첨부파일을 찾을 수 없습니다.")
    p = Path(att.stored_path)
    if not p.exists():
        raise HTTPException(status_code=410, detail="저장된 파일이 사라졌습니다.")
    return att, p


@router.get("/{settlement_id}/attachments/{attachment_id}/download")
async def download_attachment(
    settlement_id: uuid.UUID,
    attachment_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    att, p = await _resolve_attachment(db, settlement_id, attachment_id, user)
    return FileResponse(
        path=str(p),
        media_type=att.content_type or "application/octet-stream",
        filename=att.filename,
    )


@router.get("/{settlement_id}/attachments/{attachment_id}/preview")
async def preview_attachment(
    settlement_id: uuid.UUID,
    attachment_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    """브라우저 인라인 표시용 — Content-Disposition: inline."""
    att, p = await _resolve_attachment(db, settlement_id, attachment_id, user)
    return FileResponse(
        path=str(p),
        media_type=att.content_type or "application/octet-stream",
        filename=att.filename,
        content_disposition_type="inline",
    )


@router.delete("/{settlement_id}/attachments/{attachment_id}", response_model=ApiResponse[dict])
async def delete_attachment(
    settlement_id: uuid.UUID,
    attachment_id: uuid.UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(current_user),
):
    await _load_settlement_for_user(db, settlement_id, user)
    row = await db.execute(
        select(SettlementAttachment).where(
            SettlementAttachment.id == str(attachment_id),
            SettlementAttachment.settlement_id == str(settlement_id),
        )
    )
    att = row.scalar_one_or_none()
    if not att:
        raise HTTPException(status_code=404, detail="첨부파일을 찾을 수 없습니다.")
    p = Path(att.stored_path)
    await db.delete(att)
    await db.commit()
    # 파일은 DB 삭제 후 best-effort — 실패해도 메타데이터 일관성은 유지됨
    try:
        p.unlink(missing_ok=True)
    except Exception:
        pass
    return ApiResponse(success=True, data={"id": str(attachment_id)})
